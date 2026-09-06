# -*- coding: utf-8 -*-
"""SOP step 4: build C2E / E2C practice sheets from vocabs.txt.

Usage:
    python _gen_practice.py build    # generate both docx files
    python _gen_practice.py verify   # re-check both files
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SCRIPTS = Path(__file__).parent
VOCAB = SCRIPTS / "vocabs.txt"

# word -> 常见汉译 (keyed by vocabs.txt spelling)
ZH = {
    "atmosphere": "大气；大气层",
    "hydrosphere": "水圈",
    "lithosphere": "岩石圈",
    "oxygen": "氧气",
    "oxide": "氧化物",
    "carbon_dioxide": "二氧化碳",
    "hydrogen": "氢气",
    "core": "地核；核心",
    "crust": "地壳；外壳",
    "mantle": "地幔；覆盖层",
    "longitude": "经度",
    "latitude": "纬度",
    "horizon": "地平线；眼界",
    "altitude": "海拔；高度",
    "disaster": "灾难",
    "mishap": "小事故；晦气事",
    "catastrophic": "灾难性的",
    "calamity": "大灾祸；不幸事件",
    "endanger": "危及；使濒危",
    "jeopardize": "危及；损害",
    "destructive": "破坏性的；毁灭性的",
    "El_Nino": "厄尔尼诺现象",
    "greenhouse": "温室",
    "phenomenon": "现象",
    "pebble": "卵石；小圆石",
    "magnet": "磁铁；磁体",
    "ore": "矿石",
    "mineral": "矿物；矿物质的",
    "marble": "大理石",
    "quartz": "石英",
    "granite": "花岗岩",
    "gust": "一阵狂风；阵风",
    "breeze": "微风；轻风",
    "monsoon": "季风；雨季",
    "gale": "强风；大风",
    "hurricane": "飓风",
    "tornado": "龙卷风",
    "typhoon": "台风",
    "volcano": "火山",
    "erupt": "（火山）喷发；爆发",
    "magma": "岩浆",
    "thermodynamic": "热力学的",
    "smog": "烟雾；雾霾",
    "fume": "（有害的）烟；气",
    "mist": "薄雾；雾气",
    "tsunami": "海啸",
    "drought": "旱灾；干旱",
    "flooding": "洪涝；泛滥",
    "torrent": "湍流；急流",
    "earthquake": "地震",
    "seismic": "地震的",
    "avalanche": "雪崩；山崩",
    "terrain": "地形；地势",
    "landscape": "风景；地貌",
    "continent": "大陆；洲",
    "cave": "洞穴；山洞",
    "cliff": "悬崖；峭壁",
    "glacier": "冰川",
    "swamp": "沼泽",
    "delta": "三角洲",
    "plain": "平原",
    "plateau": "高原",
    "oasis": "绿洲",
    "globe": "地球；球体",
    "hemisphere": "半球",
    "equator": "赤道",
    "Arctic": "北极（的）",
    "Antarctic": "南极（的）",
    "pole": "地极；磁极",
    "polar": "极地的",
    "axis": "轴；地轴",
    "deteriorate": "恶化；变坏",
    "aggravate": "使加重；使恶化",
    "degrade": "使降级；使退化",
    "upgrade": "使升级；改善",
    "erode": "侵蚀；腐蚀",
    "Mediterranean": "地中海",
    "Atlantic": "大西洋",
    "Pacific": "太平洋",
    "ocean": "海洋",
    "marine": "海的；海产的",
    "navigation": "航行；导航",
    "gulf": "海湾；鸿沟",
    "beach": "海滩；沙滩",
    "coast": "海岸；海滨",
    "shore": "（海、湖）岸；滨",
    "tide": "潮；潮汐",
    "current": "洋流；水流",
    "brook": "小溪",
    "stream": "小河；溪流",
    "source": "源头；来源",
    "shallow": "浅的；浅薄的",
    "superficial": "表面的；肤浅的",
    "flat": "平坦的；扁平的",
    "smooth": "平滑的；光滑的",
    "rough": "粗糙的；崎岖的",
    "sandy": "多沙的；沙质的",
    "stony": "多石的；石质的",
    "vertical": "垂直的；竖直的",
    "steep": "陡的；陡峭的",
    "parallel": "平行的；类似的",
    "narrow": "狭窄的；有限的",
    "Oceania": "大洋洲",
    "mainland": "大陆；本土",
    "peninsula": "半岛",
    "climate": "气候",
    "weather": "天气",
    "meteorology": "气象学",
    "mild": "温和的；温润的",
    "heating": "暖气；供暖；加热",
    "moderate": "温和的；适度的",
    "warm": "温暖的；暖和的",
    "thermal": "热的；保暖的",
    "tropics": "热带（地区）",
    "arid": "干旱的；干燥的",
    "moist": "湿润的；潮湿的",
    "damp": "潮湿的；微湿的",
    "humid": "（空气）潮湿的；湿润的",
    "snowy": "多雪的；下雪的",
    "frost": "霜；霜冻",
    "hail": "冰雹；下冰雹",
    "thaw": "（使）解冻；融化",
    "chill": "寒意；使变冷",
    "freeze": "（使）结冰；冻结",
    "frigid": "严寒的；寒冷的",
    "tremble": "颤抖；哆嗦",
    "shiver": "（因冷、惧）发抖",
    "thunder": "雷（声）；打雷",
    "lightning": "闪电",
    "stormy": "有暴风雨的；激烈的",
    "downpour": "倾盆大雨",
    "rainfall": "降雨；降雨量",
    "sprinkle": "洒；下稀疏小雨",
    "rainbow": "彩虹",
    "shower": "阵雨；淋浴",
    "Celsius": "摄氏（的）",
    "temperature": "温度；气温",
    "forecast": "（天气）预报；预测",
    "peak": "山顶；顶峰",
    "mount": "山峰（用于山名）",
    "mountain": "山；高山",
    "range": "山脉；范围",
    "ridge": "山脊；分水岭",
    "slope": "斜坡；山坡",
    "valley": "山谷；流域",
    "hillside": "山腰；山坡",
    "overlook": "俯瞰；眺望",
    "southern": "南方的；南部的",
    "southeast": "东南（方）",
    "southwest": "西南（方）",
    "northeast": "东北（方）",
    "northwest": "西北（方）",
    "eastern": "东部的；东方的",
    "oriental": "东方的（尤指风格）",
    "inevitable": "不可避免的",
    "irreversible": "不可逆的；不可挽回的",
    "irregularly": "不规则地；无规律地",
    "inappropriate": "不恰当的；不适宜的",
    "abnormal": "反常的；异常的",
    "sediment": "沉积物；沉淀物",
    "silt": "淤泥；泥沙",
    "muddy": "泥泞的；浑浊的",
    "clay": "黏土；泥土",
    "dirt": "尘土；污物",
    "rural": "乡村的；乡下的",
    "suburb": "郊区；近郊",
    "outskirts": "市郊；郊外",
    "remote": "偏远的；遥远的",
    "desolate": "荒凉的；无人烟的",
    "distant": "遥远的；远隔的",
    "adjacent": "毗邻的；邻近的",
    "toxic": "有毒的",
    "pollution": "污染",
    "pollutant": "污染物",
    "contaminate": "污染；弄脏",
    "geology": "地质学",
    "border": "边界；国界",
    "margin": "边缘；页边空白",
    "fringe": "边缘；外围",
    "plate": "板块；板；盘子",
    "debris": "碎片；残骸",
    "crack": "裂缝；裂纹",
    "gap": "缺口；间隙；差距",
    "splendid": "辉煌的；极好的",
    "grand": "宏伟的；壮观的",
    "magnificent": "壮丽的；宏伟的",
    "super": "超级的；顶好的",
    "interesting": "有趣的",
    "dramatic": "戏剧性的；急剧的",
    "wilderness": "荒野；荒原",
    "desert": "沙漠；荒漠",
    "deforest": "滥伐（森林）",
    "barren": "贫瘠的；荒芜的",
    "fertile": "肥沃的；多产的",
    "fertilize": "施肥；使肥沃",
    "solar": "太阳的；太阳能的",
    "lunar": "月亮的；农历的",
    "calendar": "日历；历法",
    "sunrise": "日出（时分）",
    "sunset": "日落（时分）",
    "eclipse": "（日、月）食",
    "dusk": "黄昏；薄暮",
    "heaven": "天堂；天空",
    "paradise": "天堂；乐园",
    "sunshine": "阳光；日照",
    "shade": "荫；树荫",
    "shadow": "影子；阴影",
    "vapor": "水汽；蒸汽",
    "evaporate": "（使）蒸发",
    "circulate": "（使）循环；流通",
    "precipitate": "（使）沉淀；凝结",
    "reservoir": "水库；蓄水池",
    "waterfall": "瀑布",
    "fountain": "喷泉",
    "spring": "泉；泉水",
    "dew": "露水",
    "pour": "倾倒；倒灌",
    "drain": "排水；下水道",
    "drip": "滴落；水滴",
    "drown": "（使）溺水；淹没",
    "blow": "吹；刮",
    "puff": "喷出（烟、气）；一阵",
    "gush": "涌出；喷出",
    "dense": "稠密的；浓的",
    "intensity": "强度；烈度",
    "intensive": "密集的；强化的",
    "emerge": "浮现；出现",
    "flash": "闪光；闪现",
    "float": "漂浮；浮动",
    "environment": "环境",
    "surrounding": "周围的（复数：环境）",
    "condition": "条件；状况",
    "situation": "形势；处境",
    "nature": "大自然；本性",
    "natural": "自然的；天然的",
    "artificial": "人造的；人工的",
    "synthetic": "合成的；人造的",
    "petrol": "汽油（英式）",
    "gas": "气体；煤气；汽油",
    "gasoline": "汽油（美式）",
    "petroleum": "石油",
    "photosynthesis": "光合作用",
    "respire": "呼吸（作用）",
    "dioxide": "二氧化物",
    "vegetation": "植被",
    "herb": "草本植物；药草",
    "perennial": "多年生的；多年生植物",
    "botany": "植物学",
    "ecology": "生态学",
    "ecosystem": "生态系统",
    "eco-friendly": "环保的；生态友好的",
    "horticulture": "园艺（学）",
    "organism": "生物体；有机体",
    "genetics": "遗传学",
    "mutation": "突变；变异",
    "variation": "变异；差异",
    "diversity": "多样性",
    "hybridization": "杂交",
    "classify": "分类；归类",
    "reproduce": "繁殖；复制",
    "evolve": "进化；演变",
    "fluctuate": "波动；起伏",
    "reclaim": "开垦（荒地）；回收",
    "cultivate": "耕种；栽培",
    "sow": "播种",
    "harvest": "收割；收成",
    "pluck": "采摘；拔",
    "pick": "采摘；挑选",
    "yield": "产量；出产",
    "rear": "饲养；培育",
    "arable": "可耕的；耕地",
    "plough": "犁；耕地",
    "spade": "铁锹；铲",
    "rake": "耙；耙地",
    "stack": "堆叠；一堆",
    "heap": "堆；堆积",
    "bundle": "捆；束",
    "bunch": "束；串",
    "vase": "花瓶",
    "sunlight": "阳光",
    "short-day": "短日照的",
    "shade-tolerant": "耐阴的",
    "fungus": "真菌",
    "mold": "霉菌；霉",
    "pollen": "花粉",
    "germinate": "发芽；萌发",
    "seed": "种子；播种",
    "burgeon": "萌芽；迅速生长",
    "bud": "芽；花蕾",
    "flower": "花；开花",
    "blossom": "（果树的）花；开花",
    "bloom": "开花；盛开花朵",
    "scent": "气味；香味",
    "aromatic": "芳香的",
    "ripen": "（使）成熟",
}

LABEL = {"C2E": "汉译英", "E2C": "英译汉"}


def parse_vocab():
    words = []
    for ln in VOCAB.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^([A-Za-z][A-Za-z_-]*?)\s+\d+\s*$", ln.strip())
        if m:
            words.append(m.group(1))
    return words


def setup_page(sec):
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.0)
    sec.bottom_margin = Cm(1.0)
    sec.left_margin = Cm(1.0)
    sec.right_margin = Cm(1.0)
    sec.footer_distance = Cm(0.5)


def set_cols(sec, num, space_twips=240):
    cols = sec._sectPr.find(qn("w:cols"))
    if cols is None:
        from docx.oxml import OxmlElement
        cols = OxmlElement("w:cols")
        sec._sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


def base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(8.5)
    rf = st.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "宋体")


def build(kind):
    words = parse_vocab()
    assert set(words) == set(ZH), (
        f"vocab/dict mismatch: missing={set(words) - set(ZH)} extra={set(ZH) - set(words)}")

    doc = Document()
    base_style(doc)
    sec0 = doc.sections[0]
    setup_page(sec0)

    t = doc.add_paragraph()
    r = t.add_run(f"{kind} {LABEL[kind]}练习（共 {len(words)} 词，答案见页脚）")
    r.bold = True
    r.font.size = Pt(9.5)
    t.paragraph_format.space_after = Pt(4)
    t.paragraph_format.space_before = Pt(0)

    sec1 = doc.add_section(WD_SECTION_START.CONTINUOUS)
    setup_page(sec1)
    set_cols(sec1, 3)

    for i, w in enumerate(words, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.55)
        pf.first_line_indent = Cm(-0.55)
        disp = w.replace("_", " ")
        if kind == "C2E":
            n = min(24, max(10, len(disp) * 2))
            p.add_run(f"{i}. ")
            p.add_run("_" * n)
            p.add_run(" " + ZH[w])
        else:
            p.add_run(f"{i}. {disp} ")
            p.add_run("_" * 18)

    footer = sec0.footer
    fp = footer.paragraphs[0]
    parts = []
    for i, w in enumerate(words, 1):
        if kind == "C2E":
            parts.append(f"{i}.{w.replace('_', ' ')}")
        else:
            parts.append(f"{i}.{ZH[w]}")
    fp.text = "答案：" + "  ".join(parts)
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fp.paragraph_format.line_spacing = 1.0
    for run in fp.runs:
        run.font.size = Pt(4.5)

    out = SCRIPTS / f"{kind}.docx"
    doc.save(out)
    print(f"[build] {out.name}: words={len(words)} entries_written={len(words)}")


def verify():
    words = parse_vocab()
    ok = True
    for kind in ("C2E", "E2C"):
        doc = Document(SCRIPTS / f"{kind}.docx")
        secs = doc.sections
        cols = secs[-1]._sectPr.find(qn("w:cols"))
        ncols = cols.get(qn("w:num")) if cols is not None else None
        entry_paras = [p.text.strip() for p in doc.paragraphs
                       if re.match(r"^\d+\.\s", p.text.strip())]
        nums = [int(re.match(r"^(\d+)\.", t).group(1)) for t in entry_paras]
        seq_ok = nums == list(range(1, len(words) + 1))
        footer = "\n".join(p.text for s in secs for p in s.footer.paragraphs)
        if kind == "E2C":
            body = "\n".join(doc.paragraphs[i].text for i in range(len(doc.paragraphs)))
            missing = [w for w in words
                       if not re.search(r"(?i)(?<![A-Za-z])" + re.escape(w.replace("_", " ")) + r"(?![A-Za-z])", body)]
            bad_zh = [w for w in words if ZH[w] not in footer]
        else:
            missing = [w for w in words if w.replace("_", " ") not in footer]
            bad_zh = [t for t, w in zip(entry_paras, words) if ZH[w] not in t]
        line = (f"[{kind}] sections={len(secs)} cols={ncols} entries={len(entry_paras)} "
                f"seq_ok={seq_ok} missing={missing if missing else 'NONE'} "
                f"zh_bad={bad_zh if bad_zh else 'NONE'}")
        print(line)
        ok = ok and len(secs) == 2 and ncols == "3" and seq_ok and not missing and not bad_zh
    print("OVERALL:", "PASS" if ok else "FAIL")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "build"
    if cmd == "build":
        build("C2E")
        build("E2C")
    elif cmd == "verify":
        verify()
    else:
        print("unknown command")
        sys.exit(1)
