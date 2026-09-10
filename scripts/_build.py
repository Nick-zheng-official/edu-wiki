# -*- coding: utf-8 -*-
# Temporary mechanical builder (idempotent): passages.docx, C2E.docx, E2C.docx.
# Content is hand-authored in _article4.txt / _vocab_data.tsv. Safe to delete.
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

BASE = r"d:\edu-wiki\scripts\\"

def style_run(run, size, bold=False):
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), 'SimSun')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')

def set_normal(doc, size):
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), 'SimSun')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')

def set_three_cols(section):
    cols = section._sectPr.xpath('./w:cols')[0]
    cols.set(qn('w:num'), '3')
    cols.set(qn('w:space'), '397')

def add_footer_columns(section, answers, per_seg=1):
    """Page-bottom answer key: tiny font, 3 tab-stop columns."""
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.text = ""
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.tab_stops.add_tab_stop(Cm(6.4), WD_TAB_ALIGNMENT.LEFT)
    pf.tab_stops.add_tab_stop(Cm(12.8), WD_TAB_ALIGNMENT.LEFT)
    # group answers: per_seg answers per column cell, 3 segments per row
    rows = []
    items_per_row = 3 * per_seg
    for i in range(0, len(answers), items_per_row):
        row = answers[i:i + items_per_row]
        segs = []
        for c in range(3):
            seg = "  ".join(row[c * per_seg:(c + 1) * per_seg])
            segs.append(seg)
        rows.append("\t".join(segs))
    run = p.add_run("\n".join(rows))
    style_run(run, 7)

def add_entry(doc, runs_text, size=10.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    for t, b in runs_text:
        style_run(p.add_run(t), size, bold=b)

# ---------- 1. passages.docx ----------
art_lines = [l.strip() for l in open(BASE + "_article4.txt", encoding="utf-8").read().splitlines() if l.strip()]
title, paras = art_lines[0], art_lines[1:]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # Letter
sec.left_margin = sec.right_margin = Cm(3.2)
sec.top_margin = sec.bottom_margin = Cm(2.5)
set_normal(doc, 11)

h = doc.add_paragraph(style='Heading 1')
style_run(h.add_run(title), 18, bold=True)
for para in paras:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.first_line_indent = Cm(0.74)
    style_run(p.add_run(para), 11)
doc.save(BASE + "passages.docx")
print("passages.docx done:", title)

# ---------- load vocab data ----------
data = []
for line in open(BASE + "_vocab_data.tsv", encoding="utf-8"):
    line = line.rstrip("\n")
    if not line.strip():
        continue
    word, spec, zh = line.split("\t")
    forms = [tuple(x.split(":")) for x in spec.split(",")] if spec else []
    data.append((word.strip(), forms, zh.strip()))

master = []
for line in open(BASE + "vocabs.txt", encoding="utf-8"):
    parts = line.split()
    if len(parts) == 2 and parts[0][0].isalpha():
        master.append(parts[0].replace("_", " "))
tsv_words = [w for w, _, _ in data]
assert [w for w in master if w not in tsv_words] == []
assert [w for w in tsv_words if w not in master] == []
print("vocab data rows:", len(data))

# ---------- 2 & 3. worksheets ----------
def build_worksheet(fname, kind, chunk_size, bottom, per_seg):
    doc = Document()
    set_normal(doc, 10.5)
    chunks = [data[i:i + chunk_size] for i in range(0, len(data), chunk_size)]
    answer_pages = []
    for ci, chunk in enumerate(chunks):
        if ci > 0:
            doc.add_section(WD_SECTION.NEW_PAGE)
        answers = []
        for word, forms, zh in chunk:
            n = master.index(word) + 1
            if kind == "C2E":
                runs = [(f"{n}. ", False), (zh + " ", False)]
                for pos, _ans in forms:
                    runs.append((f"{pos}.________ ", False))
                add_entry(doc, runs)
                answers.append(f"{n}." + ",".join(f"{pos} {ans}" for pos, ans in forms))
            else:
                add_entry(doc, [(f"{n}. {word} ", False), ("____________", False)])
                answers.append(f"{n}.{zh}")
        answer_pages.append(answers)

    for si, section in enumerate(doc.sections):
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(1.0)
        section.top_margin = Cm(1.0)
        section.bottom_margin = Cm(bottom)
        section.footer_distance = Cm(0.4)
        section.header_distance = Cm(0.5)
        set_three_cols(section)
        add_footer_columns(section, answer_pages[si], per_seg=per_seg)

    doc.save(BASE + fname)
    print(fname, "done: pages =", len(chunks))

build_worksheet("C2E.docx", "C2E", 40, bottom=5.0, per_seg=1)
build_worksheet("E2C.docx", "E2C", 60, bottom=4.0, per_seg=2)
