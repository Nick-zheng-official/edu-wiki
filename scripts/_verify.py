# -*- coding: utf-8 -*-
# Temporary mechanical verifier for generated docx. Safe to delete.
import re
from docx import Document

BASE = r"d:\edu-wiki\scripts\\"
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# ---- passages.docx ----
doc = Document(BASE + "passages.docx")
paras = [p.text for p in doc.paragraphs if p.text.strip()]
art = open(BASE + "_article4.txt", encoding="utf-8").read()
art_paras = [l.strip() for l in art.split("\n\n") if l.strip()]
print("=== passages.docx ===")
print("title:", paras[0])
print("paragraph count:", len(paras), " expected:", len(art_paras))
ok = all(a == b for a, b in zip(paras, art_paras)) and len(paras) == len(art_paras)
print("content matches _article4.txt:", ok)

# ---- data ----
data = []
for line in open(BASE + "_vocab_data.tsv", encoding="utf-8"):
    line = line.rstrip("\n")
    if not line.strip():
        continue
    word, spec, zh = line.split("\t")
    forms = [tuple(x.split(":")) for x in spec.split(",")]
    data.append((word.strip(), forms, zh.strip()))

for fname, chunk in [("C2E.docx", 40), ("E2C.docx", 60)]:
    doc = Document(BASE + fname)
    print(f"\n=== {fname} ===")
    print("sections:", len(doc.sections), " expected:", -(-len(data)//chunk))
    for i, s in enumerate(doc.sections):
        cols = s._sectPr.xpath('./w:cols')[0]
        num = cols.get(W + 'num')
        linked = s.footer.is_linked_to_previous
        ftext = s.footer.paragraphs[0].text
        print(f"  page {i+1}: cols={num} linked={linked} footer_len={len(ftext)}")
    texts = [p.text for p in doc.paragraphs]
    body_text = "\n".join(texts)
    # count entries by leading number pattern
    nums = re.findall(r'(?m)^(\d+)\.', body_text)
    print("numbered entries:", len(nums), " first/last:", nums[0], nums[-1])
    # every master word present in body
    missing = [w for w, _, _ in data if w not in body_text]
    print("headwords missing from body:", missing)
    # footer checks: answers cover every entry
    for i, s in enumerate(doc.sections):
        ftext = s.footer.paragraphs[0].text
        chunk_items = data[i*chunk:(i+1)*chunk]
        start_n = i*chunk + 1
        for j, (word, forms, zh) in enumerate(chunk_items):
            n = start_n + j
            tag = f"{n}."
            if tag not in ftext:
                print("  ANSWER TAG MISSING:", tag, word)
        # C2E: every English answer appears in footer
        if fname == "C2E.docx":
            for word, forms, zh in chunk_items:
                for pos, ans in forms:
                    if ans not in ftext:
                        print("  ANSWER WORD MISSING:", ans, "(", word, ")")
        else:
            for word, forms, zh in chunk_items:
                head_zh = zh.split("；")[0]
                if head_zh not in ftext:
                    print("  ZH ANSWER MISSING:", head_zh, "(", word, ")")
    # sample entries
    print("sample entries:")
    for t in texts[:3]:
        print("   ", t[:90])
    print("    ...")
    for t in texts[-2:]:
        print("   ", t[:90])
    print("sample footer (page1, 200 chars):", doc.sections[0].footer.paragraphs[0].text[:200])
