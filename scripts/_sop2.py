import re
import sys
from pathlib import Path

from docx import Document

SCRIPTS = Path(__file__).parent
VOCAB = SCRIPTS / "export-vocabs.txt"
DOCX = SCRIPTS / "passages.docx"
B3_FILE = SCRIPTS / "_b3_words.txt"

B1 = ["sunshine","shade","shadow","vapor","evaporate","circulate","precipitate","reservoir","waterfall","fountain","spring","dew","pour","drain","drip","drown","blow","puff","gush","dense","intensity","intensive","emerge","flash","float","environment","surrounding","nature","natural","mist","brook","stream","source","torrent","oasis","swamp","delta","glacier","thaw","frost","hail","freeze","moist","damp","humid","arid","drought","flooding","shallow","tide","current","sprinkle","rainfall","downpour","monsoon","stormy","thunder","lightning","hurricane","tornado","breeze","tsunami","mild","moderate","snowy"]

B2 = ["lithosphere","core","crust","mantle","magma","volcano","erupt","earthquake","seismic","tremble","shiver","avalanche","terrain","landscape","continent","cave","cliff","plain","plateau","mainland","peninsula","deteriorate","aggravate","degrade","upgrade","erode","horizon","altitude","disaster","mishap","catastrophic","calamity","endanger","jeopardize","destructive","phenomenon","pebble","magnet","ore","mineral","marble","quartz","granite","globe","hemisphere","longitude","latitude","equator","pole","polar","axis","Arctic","Antarctic","thermodynamic","gale","gust","smooth","rough","steep","vertical","narrow","parallel","flat","stony","sandy","superficial","warm","thermal","chill","smog","fume","frigid"]

ARTICLES = {1: "_article1.txt", 2: "_article2.txt", 3: "_article3.txt"}


def current_words():
    return [ln.strip() for ln in VOCAB.read_text(encoding="utf-8").splitlines() if ln.strip()]


def remove_batch(n):
    batch = {w.lower() for w in (B1 if n == 1 else B2)}
    words = current_words()
    kept, removed = [], []
    for w in words:
        (removed if w.lower() in batch else kept).append(w)
    not_found = batch - {w.lower() for w in removed}
    VOCAB.write_text("\n".join(kept) + ("\n" if kept else ""), encoding="utf-8")
    print(f"[batch{n}] file_had={len(words)} removed={len(removed)} kept={len(kept)} not_found={sorted(not_found)}")


def clear_and_capture():
    b3 = current_words()
    (SCRIPTS / "_b3_words.txt").write_text("\n".join(b3), encoding="utf-8")
    VOCAB.write_text("", encoding="utf-8")
    print(f"[clear] captured batch3={len(b3)} words; export-vocabs.txt cleared")


def article_body(n):
    text = (SCRIPTS / ARTICLES[n]).read_text(encoding="utf-8").strip()
    parts = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return parts[0], parts[1:]


def batch3_words():
    return [ln.strip() for ln in B3_FILE.read_text(encoding="utf-8").splitlines() if ln.strip()]


def _write_article(n, mode):
    title, paras = article_body(n)
    wc = sum(len(p.split()) for p in paras)
    if not (600 <= wc <= 1000):
        print(f"[{mode}{n}] REFUSED: body_words={wc} outside 600-1000")
        sys.exit(2)
    doc = Document(DOCX)
    if mode == "replace":
        heads = [i for i, p in enumerate(doc.paragraphs) if p.style.name.startswith("Heading")]
        idx = heads[n - 1]
        for p in list(doc.paragraphs[idx:]):
            p._element.getparent().remove(p._element)
    doc.add_heading(title, level=1)
    for p in paras:
        doc.add_paragraph(p)
    doc.save(DOCX)
    print(f"[{mode}{n}] '{title[:60]}...' paragraphs={len(paras)} body_words={wc}")


def verify():
    b3 = batch3_words()
    all_words = B1 + B2 + b3
    dupes = {w for w in all_words if all_words.count(w) > 1}
    doc = Document(DOCX)
    paras = doc.paragraphs
    head_idx = [i for i, p in enumerate(paras) if p.style.name.startswith("Heading")]
    print("all headings:", [paras[i].text[:45] for i in head_idx])
    batches = {1: B1, 2: B2, 3: b3}
    ok = True
    for n in (1, 2, 3):
        title = article_body(n)[0]
        idx = next(i for i in head_idx if paras[i].text.strip() == title)
        end = next((i for i in head_idx if i > idx), len(paras))
        body = "\n".join(p.text for p in paras[idx + 1:end])
        wc = len(body.split())
        missing = [w for w in batches[n] if not re.search(r"(?i)(?<![a-z])" + re.escape(w), body)]
        print(f"[article{n}] '{title[:45]}...' body_words={wc} in_range={'YES' if 600 <= wc <= 1000 else 'NO'} "
              f"batch_words={len(batches[n])} missing={missing if missing else 'NONE'}")
        ok = ok and 600 <= wc <= 1000 and not missing
    full = "\n".join(p.text for p in paras)
    doc_missing = [w for w in all_words if not re.search(r"(?i)(?<![a-z])" + re.escape(w), full)]
    print("doc-level missing:", doc_missing if doc_missing else "NONE")
    print("dupes across batches:", dupes if dupes else "NONE")
    print("vocab file now:", repr(VOCAB.read_text(encoding='utf-8')))
    print("OVERALL:", "PASS" if ok and not doc_missing and not dupes else "FAIL")


cmd = sys.argv[1]
if cmd == "remove1":
    remove_batch(1)
elif cmd == "append1":
    _write_article(1, "append")
elif cmd == "remove2":
    remove_batch(2)
elif cmd == "append2":
    _write_article(2, "append")
elif cmd == "clear3":
    clear_and_capture()
elif cmd == "append3":
    _write_article(3, "append")
elif cmd == "replace3":
    _write_article(3, "replace")
elif cmd == "verify":
    verify()
else:
    print("unknown command")
    sys.exit(1)
