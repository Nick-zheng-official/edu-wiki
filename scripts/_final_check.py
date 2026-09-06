# -*- coding: utf-8 -*-
"""SOP step 5: final overall check of every generated artifact."""
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

SCRIPTS = Path(__file__).parent


def check_passages():
    doc = Document(SCRIPTS / "passages.docx")
    ok = True
    for i, p in enumerate(doc.paragraphs):
        if p.style.name.startswith("Heading"):
            r = p.runs[0] if p.runs else None
            name = r.font.name if r else None
            color = r.font.color.rgb if r and r.font.color and r.font.color.rgb else None
            bold = r.bold if r else None
            ppr = p._element.find(qn("w:pPr"))
            has_break = False
            if ppr is not None:
                for pb in ppr.iter(qn("w:pageBreakBefore")):
                    has_break = True
            print(f"[heading] '{p.text[:40]}' font={name} color={color} bold={bold} "
                  f"page_break={has_break}")
            ok = ok and name == "Times New Roman" and bold and str(color) == "000000"
    print("[passages] headings_ok:", ok)


def check_vocabs_files():
    exp = (SCRIPTS / "export-vocabs.txt").read_text(encoding="utf-8").strip()
    print("[export-vocabs] empty:", exp == "")
    vw = [ln for ln in (SCRIPTS / "vocabs.txt").read_text(encoding="utf-8").splitlines()
          if re.match(r"^[A-Za-z][A-Za-z_-]*\s+\d+\s*$", ln.strip())]
    print(f"[vocabs.txt] words={len(vw)}")
    return len(vw)


def check_c2e_e2c(n_words):
    for kind in ("C2E", "E2C"):
        doc = Document(SCRIPTS / f"{kind}.docx")
        entry_paras = [p for p in doc.paragraphs if re.match(r"^\d+\.\s", p.text.strip())]
        footer = "\n".join(p.text for s in doc.sections for p in s.footer.paragraphs)
        ans_count = len(re.findall(r"(?:^|\s)\d+\.", footer.replace("答案：", "答案： ")))
        print(f"[{kind}] entries={len(entry_paras)} footer_answers~={ans_count} "
              f"footer_chars={len(footer)}")


if __name__ == "__main__":
    check_passages()
    n = check_vocabs_files()
    check_c2e_e2c(n)
